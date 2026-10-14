"""Excel to PDF Handler"""

import csv
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import PageBreak, SimpleDocTemplate, Table, TableStyle


def _read_sheets(input_path: Path) -> List[List]:
    ext = input_path.suffix.lower()
    if ext == ".csv":
        with open(input_path, "r", encoding="utf-8", errors="replace", newline="") as handle:
            rows = [["" if cell is None else str(cell) for cell in row] for row in csv.reader(handle)]
        return [[input_path.stem or "Sheet1", rows]]

    workbook = openpyxl.load_workbook(input_path, data_only=True)
    sheets = []
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        sheets.append([name, rows])
    workbook.close()
    return sheets


def _normalize(rows: List[List[str]]) -> List[List[str]]:
    if not rows:
        return [[""]]
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def handler(job_dir: Path, inputs: List[Dict], options: Dict[str, Any]) -> Dict:
    """Convert a simple Excel/CSV table into a PDF."""
    orientation = options.get("orientation", "auto")
    sheets_opt = options.get("sheets", "all")

    input_path = Path(inputs[0]["path"])
    sheets = _read_sheets(input_path)

    if sheets_opt == "first":
        sheets = sheets[:1]

    widest = max((max((len(row) for row in rows), default=0) for _, rows in sheets), default=0)
    if orientation == "wide" or (orientation == "auto" and widest > 6):
        page_size = landscape(A4)
    else:
        page_size = A4

    output_path = job_dir / "out/output.pdf"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    story = []
    for index, (name, rows) in enumerate(sheets):
        data = _normalize(rows)
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CFCFC8")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4F3EF")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(table)
        if index < len(sheets) - 1:
            story.append(PageBreak())

    document = SimpleDocTemplate(str(output_path), pagesize=page_size)
    document.build(story)

    return {
        "ok": True,
        "outputs": [
            {"path": "out/output.pdf", "name": "output.pdf", "mime": "application/pdf"}
        ],
        "stats": {
            "total_sheets": len(sheets),
            "orientation": "wide" if page_size is not A4 else "tall",
            "sheets_processed": len(sheets),
        },
        "progress": 100,
        "note": "BETA: Charts and pictures not included",
    }
