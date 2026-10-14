"""PDF to Excel Handler.

Extracts real tables from a PDF using pdfplumber and writes them to an
``.xlsx`` workbook. Pages without detectable tables fall back to their plain
text lines so the output is still useful. Scans (images without a text layer)
are not supported.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pdfplumber
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_PAGES = 100
_INVALID_SHEET_CHARS = set('[]:*?/\\')

_HEADER_FILL = PatternFill("solid", fgColor="F4F3EF")
_HEADER_FONT = Font(bold=True)
_THIN = Side(style="thin", color="CFCFC8")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").strip()


def _normalize(table: List[List[Any]]) -> List[List[str]]:
    width = max((len(row) for row in table), default=0)
    if width == 0:
        return []
    return [[_clean(cell) for cell in row] + [""] * (width - len(row)) for row in table]


def _sheet_title(name: str, used: set) -> str:
    title = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in name)[:31] or "Sheet"
    candidate = title
    index = 2
    while candidate in used:
        suffix = f" ({index})"
        candidate = title[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def _page_rows(page: pdfplumber.page.Page) -> Tuple[List[List[str]], int]:
    """Extract rows for one page. Returns (rows, table_count)."""
    rows: List[List[str]] = []
    tables = 0
    for table in page.extract_tables() or []:
        normalized = _normalize(table)
        if not normalized:
            continue
        if rows:
            rows.append([])
        rows.extend(normalized)
        tables += 1
    if rows:
        return rows, tables

    text = page.extract_text() or ""
    lines = [line for line in text.splitlines() if line.strip()]
    return [[line] for line in lines], 0


def _style_sheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    first = None
    for row in ws.iter_rows():
        if any(cell.value not in (None, "") for cell in row):
            first = row
            break
    if first is not None:
        for cell in first:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL

    widths: Dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            widths[idx] = max(widths.get(idx, 0), len("" if value is None else str(value)))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 60)

    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                cell.border = _BORDER

    ws.freeze_panes = "A2"


def handler(job_dir: Path, inputs: List[Dict], options: Dict[str, Any]) -> Dict:
    """Convert a PDF's tables into an Excel workbook."""
    if not inputs:
        return {"ok": False, "error": {"code": "NO_INPUTS", "message": "No input file provided"}}

    layout = str(options.get("layout", "per-page")).replace("_", "-")
    one_sheet = layout in ("one-sheet", "single", "combined")

    input_path = Path(inputs[0]["path"])
    if not input_path.exists():
        return {"ok": False, "error": {"code": "NO_INPUTS", "message": "Input file not found"}}

    source_stem = Path(inputs[0].get("name") or input_path.name).stem or "output"

    out_dir = job_dir / "out"
    out_dir.mkdir(parents=True, exist_ok=True)
    output_path = out_dir / "output.xlsx"

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    used_titles: set = set()
    total_pages = 0
    processed_pages = 0
    table_count = 0
    pages_with_data = 0

    try:
        with pdfplumber.open(str(input_path)) as pdf:
            total_pages = len(pdf.pages)
            page_results: List[Tuple[int, List[List[str]]]] = []
            for page_no, page in enumerate(pdf.pages[:MAX_PAGES], start=1):
                rows, tables = _page_rows(page)
                table_count += tables
                processed_pages += 1
                if rows:
                    page_results.append((page_no, rows))

            pages_with_data = len(page_results)

            if one_sheet:
                combined: List[List[str]] = []
                for _, rows in page_results:
                    if combined:
                        combined.append([])
                    combined.extend(rows)
                if combined:
                    ws = workbook.create_sheet(_sheet_title("Tables", used_titles))
                    for row in combined:
                        ws.append(row)
                    _style_sheet(ws)
            else:
                for page_no, rows in page_results:
                    ws = workbook.create_sheet(_sheet_title(f"Page {page_no}", used_titles))
                    for row in rows:
                        ws.append(row)
                    _style_sheet(ws)
    except Exception as exc:  # noqa: BLE001 - report a friendly error to the client
        return {
            "ok": False,
            "error": {
                "code": "PDF_UNREADABLE",
                "message": "Could not read this PDF. It may be encrypted or damaged.",
                "detail": str(exc),
            },
        }

    if not workbook.sheetnames:
        ws = workbook.create_sheet("No tables")
        ws["A1"] = "No tables or text were found. This looks like a scanned PDF."
        ws["A2"] = "Scans are not supported by this tool."
        workbook.active = 0

    workbook.save(output_path)

    return {
        "ok": True,
        "outputs": [
            {
                "path": "out/output.xlsx",
                "name": f"{source_stem}.xlsx",
                "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            }
        ],
        "stats": {
            "total_pages": total_pages,
            "pages_processed": processed_pages,
            "pages_with_content": pages_with_data,
            "tables_found": table_count,
            "sheets": len(workbook.sheetnames),
            "layout": "one-sheet" if one_sheet else "per-page",
        },
        "progress": 100,
        "note": "BETA: works with real text and tables; scans are not supported.",
    }
