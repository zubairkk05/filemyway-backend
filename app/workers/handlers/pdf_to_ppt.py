"""PDF to PPT Handler (PowerPoint conversion)"""

from pathlib import Path
from typing import Any, Dict, List
import pikepdf

try:
    import pypdfium2 as pdfium
    PYPDFIUM2_AVAILABLE = True
except ImportError:
    PYPDFIUM2_AVAILABLE = False


_QUALITY_SCALE = {"small": 1.5, "low": 1.5, "good": 2.0, "medium": 2.0, "best": 3.0, "high": 3.0}


def handler(job_dir: Path, inputs: List[Dict], options: Dict[str, Any]) -> Dict:
    """Convert a PDF into a PowerPoint deck (one picture slide per page)."""
    quality = options.get("quality", "good")

    if not PYPDFIUM2_AVAILABLE:
        return {
            "ok": False,
            "error": {"code": "MISSING_LIBRARY", "message": "pypdfium2 is required"},
        }

    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return {
            "ok": False,
            "error": {"code": "MISSING_LIBRARY", "message": "python-pptx is not installed"},
        }

    input_path = Path(inputs[0]["path"])
    out_dir = job_dir / "out"
    out_dir.mkdir(parents=True, exist_ok=True)

    scale = _QUALITY_SCALE.get(quality, 2.0)
    doc = pdfium.PdfDocument(str(input_path))
    total_pages = len(doc)
    pages_to_convert = list(range(min(total_pages, 50)))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slide_w = prs.slide_width
    slide_h = prs.slide_height

    for page_num in pages_to_convert:
        pil_image = doc[page_num].render(scale=scale).to_pil()
        tmp_png = out_dir / f"_slide_{page_num + 1}.png"
        pil_image.save(tmp_png, "PNG")

        slide = prs.slides.add_slide(blank_layout)
        img_w, img_h = pil_image.size
        ratio = min(slide_w / img_w, slide_h / img_h)
        width = int(img_w * ratio)
        height = int(img_h * ratio)
        left = int((slide_w - width) / 2)
        top = int((slide_h - height) / 2)
        slide.shapes.add_picture(str(tmp_png), left, top, width=width, height=height)

    doc.close()

    output_path = out_dir / "output.pptx"
    prs.save(output_path)

    for tmp_png in out_dir.glob("_slide_*.png"):
        tmp_png.unlink()

    return {
        "ok": True,
        "outputs": [
            {
                "path": "out/output.pptx",
                "name": "output.pptx",
                "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            }
        ],
        "stats": {
            "total_pages": total_pages,
            "slides_created": len(pages_to_convert),
            "quality": quality,
        },
        "progress": 100,
        "note": "BETA: Slides are pictures of the pages; text is not editable",
    }
